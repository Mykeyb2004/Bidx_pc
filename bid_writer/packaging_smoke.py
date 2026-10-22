"""Offline checks run inside the distributed executable, without reading API keys."""

from __future__ import annotations

import json
import tempfile
import time
import traceback
from pathlib import Path


def run_smoke_test(report_filename: str) -> int:
    started = time.monotonic()
    report: dict = {"ok": False, "checks": []}
    destination = Path(report_filename).expanduser().resolve()
    destination.parent.mkdir(parents=True, exist_ok=True)
    destination.write_text(json.dumps(report), encoding="utf-8")
    try:
        from .gui_state import (
            get_application_root_dir,
            get_cache_root_dir,
            get_log_dir,
            get_resource_root_dir,
            get_shared_config_dir,
            is_compiled,
        )

        resource_root = get_resource_root_dir()
        report.update(
            compiled=is_compiled(),
            resource_root=str(resource_root),
            data_root=str(get_application_root_dir()),
            shared_config_dir=str(get_shared_config_dir()),
            log_dir=str(get_log_dir()),
            cache_root_dir=str(get_cache_root_dir()),
        )
        for name in (
            "roles/system_gate_rules.md", "roles/通用投标角色.md", "roles/标书架构师.md",
            "roles/tender_section_boundaries.yaml", "config.example.yaml", ".env.example",
            "bid_writer/assets/brand/logo_128.png", "bid_writer/assets/icons/tabler/generate.png",
        ):
            if not (resource_root / name).is_file():
                raise RuntimeError(f"Missing bundled resource: {name}")
        report["checks"].append("resources")

        import openpyxl
        import pandas
        import pymupdf
        from docx import Document
        from python_calamine import CalamineWorkbook
        from .tender_markdown_converter import convert_tender_document
        from .tender_section_boundary_config import load_boundary_config

        load_boundary_config()
        with tempfile.TemporaryDirectory(prefix="bid-writer-documents-") as temporary:
            directory = Path(temporary)
            document = Document()
            document.add_paragraph("BidWriter packaging check")
            document.save(directory / "sample.docx")
            workbook = openpyxl.Workbook()
            workbook.active.append(["Name", "Value"])
            workbook.active.append(["check", 1])
            workbook.save(directory / "sample.xlsx")
            workbook.close()
            with CalamineWorkbook.from_path(str(directory / "sample.xlsx")) as calamine:
                if not calamine.get_sheet_by_index(0).to_python():
                    raise RuntimeError("Calamine produced no rows")
            if "check" not in pandas.DataFrame({"Name": ["check"]}).to_markdown(index=False):
                raise RuntimeError("pandas/tabulate check failed")
            with pymupdf.open() as pdf:
                page = pdf.new_page()
                page.insert_text((72, 72), "\n".join(
                    f"BidWriter packaging check paragraph {number}: readable PDF text."
                    for number in range(1, 6)
                ))
                pdf.save(directory / "sample.pdf")
            for suffix in ("docx", "xlsx", "pdf"):
                converted = convert_tender_document(directory / f"sample.{suffix}", directory / suffix)
                if not converted.blocks:
                    raise RuntimeError(f"No content from {suffix}")
                report["checks"].append(suffix)
        report["checks"].append("calamine/pandas/tabulate")
        # Preserve completed checks even if a native GUI library aborts the process.
        report["stage"] = "gui"
        destination.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")

        from .gui import ensure_tk_runtime
        from .ui_icons import get_brand_image, get_icon_image
        import ttkbootstrap

        ensure_tk_runtime()
        root = ttkbootstrap.Window(themename="litera")
        try:
            root.withdraw()
            for image in (get_brand_image(root, 128), get_icon_image(root, "generate")):
                if image is None:
                    raise RuntimeError("Tk failed to load bundled image")
            ttkbootstrap.Button(root, text="验证").pack()
            root.update()
            report["checks"].append("tk/ttkbootstrap/images")
        finally:
            root.destroy()
        report["ok"] = True
    except Exception:
        report["error"] = traceback.format_exc()
    report["elapsed_seconds"] = round(time.monotonic() - started, 3)
    destination.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    return 0 if report["ok"] else 1
